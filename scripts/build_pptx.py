#!/usr/bin/env python3
"""Genera docs/presentacion.pptx: 18 slides, formato sobrio (fondo blanco, texto negro, Calibri).

Si existe docs/uml.png (el diagrama que dibuja Franco) se usa en la slide 4.
Uso (desde la raíz del repo): python3 scripts/build_pptx.py
"""
import pathlib
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# Campos editables de portada: el usuario pidio completarlos despues.
MEMBERS = ["[Nombre y apellido de los integrantes]"]
DATE = "[Fecha de exposición]"

OUT = pathlib.Path("docs/presentacion.pptx")
UML = pathlib.Path("docs/uml.png")
CAPTURE = pathlib.Path("docs/captura.png")
BLACK = RGBColor(0, 0, 0)

SLIDES = [
    {
        "title": "Ajedrez: núcleo y adaptadores",
        "bullets": ["TPO de Ingeniería de Software, UADE"],
        "notes": "Presentación del grupo y del trabajo.",
        "cover": True,
    },
    {
        "title": "La consigna y lo que hicimos",
        "bullets": [
            "Obligatorio: tablero 8x8, seis tipos de pieza, turnos, captura, validación y jaque",
            "Opcional hecho: jaque mate, enroque, captura al paso, coronación",
            "Opcional hecho: tablas (3 reglas), deshacer y rehacer, interfaz gráfica",
            "Opcional hecho: oponente con IA en tres niveles",
            "Núcleo sin framework y probado sin infraestructura",
        ],
        "notes": "Todo lo obligatorio y todos los opcionales.",
    },
    {
        "title": "Arquitectura general",
        "bullets": [
            "nucleo: las reglas del juego",
            "puertos: lo que el núcleo pide al exterior",
            "adaptadores: consola y ventana Swing",
            "ai: el oponente, afuera del núcleo",
            "Las dependencias apuntan hacia adentro",
        ],
        "code": "adaptadores --> puertos --> nucleo\nai ------------------> nucleo\nMain / MainSwing: arman todo",
        "notes": "Regla: el núcleo no importa nada de afuera.",
    },
    {
        "title": "Diagrama de clases",
        "bullets": [
            "UML consistente con el código",
            "Una sola entrada: IChessGame",
            "Piezas con estrategia de movimiento",
            "Reglas especiales y de tablas enchufables",
        ],
        "image": "uml",
        "notes": "Mostrar el diagrama y recorrer los paquetes.",
    },
    {
        "title": "Núcleo y adaptadores",
        "bullets": [
            "La prueba: si cambia la tecnología, ¿esto cambia?",
            "Reglas del ajedrez: núcleo",
            "Consola, Swing, imágenes de las piezas: adaptadores",
            "La ventana se agregó sin tocar ChessGame",
        ],
        "notes": "El criterio de la cátedra, aplicado a nuestro caso.",
    },
    {
        "title": "Puertos",
        "bullets": [
            "Entrada: IChessGame (uno solo)",
            "Salida: IMoveListener, IGameStatusListener, IPromotionChooser",
            "Un puerto de salida por cada cosa que el núcleo pide afuera",
            "El observer se partió en dos (Interface Segregation)",
        ],
        "code": "interface IChessGame {\n  MoveOutcome tryMove(\n    Position from, Position to);\n  boolean undo();\n  boolean redo();\n  List<Move> legalMoves();\n}",
        "notes": "Por qué dos listeners y no uno.",
    },
    {
        "title": "Inyección de dependencias",
        "bullets": [
            "Nivel 4: todo entra por constructor",
            "Main y MainSwing arman las partidas de juego",
            "Se prueba con un stub y un mock",
            "Sumar una regla es sumar una línea al armado",
        ],
        "code": "new ChessGame(board,\n  moveListeners,\n  statusListeners,\n  specialMoves,\n  drawRules);",
        "notes": "Cada colaborador inyectado es un colaborador externo distinto.",
    },
    {
        "title": "Composición sobre herencia",
        "bullets": [
            "Una sola clase Piece",
            "Piece tiene una IMovementStrategy",
            "La reina compone torre y alfil",
            "Una pieza nueva es una estrategia nueva",
        ],
        "code": "class Piece {\n  PieceType type;\n  Color color;\n  IMovementStrategy movement;\n}",
        "notes": "Sin árbol de herencia de piezas.",
    },
    {
        "title": "Factory",
        "bullets": [
            "PieceFactory es un registro tipo -> estrategia",
            "PieceType es una clase de valor, no un enum",
            "BoardSetup.custom recibe tamaño y fila trasera",
            "Se agrega una pieza sin editar clases",
        ],
        "code": "IMovementStrategy custom =\n  new ChancellorMovementStrategy();\nfactory.register(\n  new PieceType(\"CHANCELLOR\"),\n  custom);",
        "notes": "El enum era un conjunto cerrado.",
    },
    {
        "title": "Command: jugadas que se pueden deshacer",
        "bullets": [
            "MoveCommand, CastlingCommand, PromotionCommand, EnPassantCommand",
            "Cada uno guarda lo necesario para revertir",
            "Historial y pila de rehacer en ChessGame",
            "Las casillas tocadas dan el derecho de enroque",
        ],
        "code": "interface ICommand {\n  void execute();\n  void undo();\n}",
        "notes": "Deshacer y rehacer salen del historial.",
    },
    {
        "title": "State",
        "bullets": [
            "Turno: WhiteTurnState y BlackTurnState",
            "Terminales: Checkmate, Stalemate y Draw",
            "No hay un booleano gameOver",
            "Deshacer vuelve al estado anterior",
        ],
        "code": "interface IGameState {\n  Color currentColor();\n  boolean isOver();\n  IGameState afterMove();\n}",
        "notes": "El patrón resuelve un problema real: los estados terminales.",
    },
    {
        "title": "Observer",
        "bullets": [
            "El núcleo avisa lo que pasa sin saber quién escucha",
            "ConsoleRenderer y SwingChessView implementan los listeners",
            "Los tests usan un mock que registra los avisos",
        ],
        "notes": "Dos adaptadores, un mismo puerto.",
    },
    {
        "title": "Reglas enchufables",
        "bullets": [
            "ISpecialMove: enroque, captura al paso, coronación",
            "IDrawRule: material, 50 movimientos, repetición",
            "ChessGame las recibe y no sabe cuáles son",
            "Regla nueva: una clase y una línea en el armado",
        ],
        "code": "interface ISpecialMove {\n  List<Position> extraDestinations(\n    IGameContext context,\n    Position from);\n  IMoveCommand commandFor(\n    IGameContext context,\n    Position from, Position to);\n}",
        "notes": "Open/Closed aplicado.",
    },
    {
        "title": "Extensibilidad",
        "bullets": [
            "Pieza nueva sin tocar clases existentes",
            "Tablero de otro tamaño: new Board(10, 10)",
            "Regla de tablas nueva sin editar ChessGame",
            "Cada una tiene un test",
        ],
        "notes": "Lo que puede pedir la defensa individual.",
    },
    {
        "title": "Tests",
        "bullets": [
            "Núcleo probado sin infraestructura",
            "Arrange, Act, Assert",
            "Stub y mock como dobles de prueba",
            "Los marcados TDD se escribieron antes del código",
            "Runner propio: ./scripts/test.sh",
        ],
        "notes": "Contar cuántos tests hay con ./scripts/test.sh.",
    },
    {
        "title": "La IA, afuera del núcleo",
        "bullets": [
            "Tres niveles: azar, captura codiciosa y minimax",
            "Juega a través de IChessGame",
            "Piensa en una partida de búsqueda propia",
            "No abre diálogos ni toca la pantalla",
        ],
        "notes": "Es un extra para jugar; no forma parte de las reglas.",
    },
    {
        "title": "Decisiones y cuándo se romperían",
        "bullets": [
            "Núcleo separado: sobra en un prototipo descartable",
            "Estrategias por pieza: una jerarquía corta alcanzaría si compartieran estado",
            "PieceType abierto: un enum alcanzaría con piezas fijas",
            "State con terminales: con 2 estados bastaba un enum",
            "Detalle completo en docs/justificacion",
        ],
        "notes": "Qué, por qué y cuándo se rompería, resumido.",
    },
    {
        "title": "Demostración",
        "bullets": [
            "Partida contra la IA en la ventana",
            "Enroque, coronación y captura al paso",
            "Deshacer y rehacer",
            "Gracias",
        ],
        "image": "capture",
        "notes": "Demo en vivo: abrir MainSwing.",
    },
]


def bulletize(paragraph):
    properties = paragraph._p.get_or_add_pPr()
    properties.set("marL", "285750")
    properties.set("indent", "-285750")
    bullet = etree.SubElement(properties, qn("a:buChar"))
    bullet.set("char", "•")


def add_text(slide, left, top, width, height, size, bold=False, mono=False):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    return box, frame


def style(paragraph, text, size, bold=False, mono=False):
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Courier New" if mono else "Calibri"
    run.font.color.rgb = BLACK


def main():
    if not MEMBERS or not DATE:
        sys.exit("Completá MEMBERS y DATE al principio del script (preguntale a Franco).")

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    for number, data in enumerate(SLIDES, 1):
        assert len(data["bullets"]) <= 6, "demasiados puntos en la slide " + str(number)
        for bullet in data["bullets"]:
            assert len(bullet) <= 100, "punto demasiado largo en la slide " + str(number) + ": " + bullet
        code = data.get("code")
        if code:
            assert len(code.splitlines()) <= 10, "demasiado código en la slide " + str(number)
            assert all(len(line) <= 50 for line in code.splitlines()), "línea de código larga en la slide " + str(number)

        slide = presentation.slides.add_slide(blank)
        slide.notes_slide.notes_text_frame.text = data["notes"]

        _, title_frame = add_text(slide, 0.7, 0.45, 11.9, 1.1, 34)
        style(title_frame.paragraphs[0], data["title"], 34, bold=True)

        has_side = bool(code) or data.get("image")
        body_width = 6.2 if has_side else 11.9
        _, body = add_text(slide, 0.7, 1.8, body_width, 5.0, 22)
        first = True
        if data.get("cover"):
            lines = data["bullets"] + ["Integrantes: " + ", ".join(MEMBERS), DATE]
            for line in lines:
                paragraph = body.paragraphs[0] if first else body.add_paragraph()
                first = False
                style(paragraph, line, 24)
                paragraph.space_after = Pt(14)
        else:
            for bullet in data["bullets"]:
                paragraph = body.paragraphs[0] if first else body.add_paragraph()
                first = False
                bulletize(paragraph)
                style(paragraph, bullet, 22)
                paragraph.space_after = Pt(12)

        if code:
            box, frame = add_text(slide, 7.2, 1.8, 5.5, 4.7, 18)
            box.fill.solid()
            box.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2)
            for index, line in enumerate(code.splitlines()):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                style(paragraph, line, 18, mono=True)
        elif data.get("image") == "uml" and UML.exists():
            slide.shapes.add_picture(str(UML), Inches(7.0), Inches(1.7), width=Inches(5.9))
        elif data.get("image") == "uml":
            _, placeholder = add_text(slide, 7.2, 2.7, 5.2, 1.5, 22)
            style(placeholder.paragraphs[0], "[Insertar UML de Franco]", 22)
        elif data.get("image") == "capture" and CAPTURE.exists():
            slide.shapes.add_picture(str(CAPTURE), Inches(8.2), Inches(1.5), height=Inches(5.6))

    presentation.save(str(OUT))
    print("escrito " + str(OUT) + " con " + str(len(SLIDES)) + " slides")


main()
