#!/usr/bin/env python3
"""Convierte un .md sencillo a .docx con formato académico: Times New Roman, negro.

Soporta títulos (#, ##, ###), párrafos, listas con "-", bloques de código entre
tres comillas invertidas, **negrita** y `código` en línea.
Uso: python3 scripts/md_to_docx.py docs/justificacion.md docs/justificacion.docx
"""
import re
import sys

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

BLACK = RGBColor(0, 0, 0)
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def style_run(run, size=12, bold=False, mono=False):
    font_name = "Courier New" if mono else "Times New Roman"
    run.font.name = font_name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(10 if mono else size)
    run.font.color.rgb = BLACK
    run.bold = bold


def add_inline(paragraph, text, size=12, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            style_run(paragraph.add_run(part[2:-2]), size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            style_run(paragraph.add_run(part[1:-1]), mono=True)
        else:
            style_run(paragraph.add_run(part), size=size, bold=bold)


def main():
    source, target = sys.argv[1], sys.argv[2]
    document = Document()
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet"):
        document.styles[style_name].font.name = "Times New Roman"
        document.styles[style_name].font.color.rgb = BLACK
    for document_style in document.styles:
        for border in list(document_style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    document.styles["Normal"].font.size = Pt(12)
    document.styles["Normal"].paragraph_format.space_after = Pt(6)
    in_code = False

    for raw in open(source, encoding="utf-8").read().splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(18)
            paragraph.paragraph_format.space_after = Pt(0)
            style_run(paragraph.add_run(line), mono=True)
            continue
        if not line.strip():
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            paragraph = document.add_paragraph(style="Title" if level == 1 else "Heading " + str(level - 1))
            add_inline(paragraph, heading.group(2), size={1: 16, 2: 14, 3: 12}[level], bold=True)
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
        else:
            paragraph = document.add_paragraph()
            add_inline(paragraph, line)

    document.save(target)
    print("escrito " + target)


main()
